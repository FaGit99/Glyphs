"""Build paper V8 as .docx (then converted to .odt/.pdf with LibreOffice).
usage: venv/bin/python build_v8.py <glyphs_dir> <fig_dir> <v5_pictures_dir> <out.docx>
  from glyphs/:  python Papers/build_v8/build_v8.py . Papers/build_v8 Papers/build_v8/report_examples out.docx"""
import csv, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

G, FIG, V5P, OUT = sys.argv[1:5]
FONT = "Liberation Serif"
doc = Document()

# ------------------------------------------------------------------ page and styles
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.4)
sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
TEXTW = 21.0 - 4.8


def set_font(style, size, bold=None, italic=None):
    style.font.name = FONT; style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


st = doc.styles
set_font(st["Normal"], 10.5)
pf = st["Normal"].paragraph_format
pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(5), 1.15
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, it, before in (("Heading 1", 12.5, False, 14), ("Heading 2", 10.5, True, 9), ("Heading 3", 10.5, True, 6)):
    set_font(st[name], size, bold=True, italic=it)
    f = st[name].paragraph_format
    f.space_before, f.space_after, f.keep_with_next = Pt(before), Pt(4), True
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name in ("Title", "List Bullet", "List Number", "Caption"):
    set_font(st[name], 10.5)

# ------------------------------------------------------------------ citations
REFS = {
 "kittler": "Kittler H, Marghoob AA, Argenziano G, Carrera C, Curiel-Lewandrowski C, Hofmann-Wellenhof R, et al. Standardization of terminology in dermoscopy/dermatoscopy: results of the third consensus conference of the International Society of Dermoscopy. J Am Acad Dermatol. 2016;74(6):1093–1106. doi:10.1016/j.jaad.2015.12.038",
 "argenziano": "Argenziano G, Fabbrocini G, Carli P, De Giorgi V, Sammarco E, Delfino M. Epiluminescence microscopy for the diagnosis of doubtful melanocytic skin lesions: comparison of the ABCD rule of dermatoscopy and a new 7-point checklist based on pattern analysis. Arch Dermatol. 1998;134(12):1563–1570. doi:10.1001/archderm.134.12.1563",
 "kawahara": "Kawahara J, Daneshvar S, Argenziano G, Hamarneh G. Seven-point checklist and skin lesion classification using multitask multimodal neural nets. IEEE J Biomed Health Inform. 2019;23(2):538–546. doi:10.1109/JBHI.2018.2824327",
 "codella": "Codella N, Rotemberg V, Tschandl P, Celebi ME, Dusza S, Gutman D, et al. Skin lesion analysis toward melanoma detection 2018: a challenge hosted by the International Skin Imaging Collaboration (ISIC). arXiv:1902.03368. 2019.",
 "ham": "Tschandl P, Rosendahl C, Kittler H. The HAM10000 dataset, a large collection of multi-source dermatoscopic images of common pigmented skin lesions. Sci Data. 2018;5:180161. doi:10.1038/sdata.2018.161",
 "bcn": "Combalia M, Codella NCF, Rotemberg V, Helba B, Vilaplana V, Reiter O, et al. BCN20000: dermoscopic lesions in the wild. arXiv:1908.02288. 2019.",
 "isic2020": "Rotemberg V, Kurtansky N, Betz-Stablein B, Caffery L, Chousakos E, Codella N, et al. A patient-centric dataset of images and metadata for identifying melanomas using clinical context. Sci Data. 2021;8:34. doi:10.1038/s41597-021-00815-z",
 "isicdata": "International Skin Imaging Collaboration. ISIC 2019 and ISIC 2020 challenge datasets. ISIC Archive; 2019–2020. doi:10.34970/2020-ds01. Available from: https://challenge.isic-archive.com/data",
 "eva": "Fang Y, Sun Q, Wang X, Huang T, Wang X, Cao Y. EVA-02: a visual representation for neon genesis. Image Vis Comput. 2024;149:105171. doi:10.1016/j.imavis.2024.105171",
 "panderm": "Yan S, Yu Z, Primiero C, Vico-Alonso C, Wang Z, Yang L, et al. A multimodal vision foundation model for clinical dermatology. Nat Med. 2025;31(8):2691–2702. doi:10.1038/s41591-025-03747-y",
 "vit": "Dosovitskiy A, Beyer L, Kolesnikov A, Weissenborn D, Zhai X, Unterthiner T, et al. An image is worth 16x16 words: transformers for image recognition at scale. In: International Conference on Learning Representations (ICLR); 2021. arXiv:2010.11929",
 "platt": "Platt J. Probabilistic outputs for support vector machines and comparisons to regularized likelihood methods. In: Smola A, Bartlett P, Schölkopf B, Schuurmans D, editors. Advances in Large Margin Classifiers. Cambridge (MA): MIT Press; 1999. p. 61–74.",
 "timm": "Wightman R. PyTorch Image Models (timm) [software]. GitHub; 2019. doi:10.5281/zenodo.4414861",
 "liopyris": "Liopyris K, Navarrete-Dechent C, Marchetti MA, Rotemberg V, Apalla Z, Argenziano G, et al. Expert agreement on the presence and spatial localization of melanocytic features in dermoscopy. J Invest Dermatol. 2024;144(3):531–539. doi:10.1016/j.jid.2023.01.045",
 "brancaccio": "Brancaccio G, Balato A, Malvehy J, Puig S, Argenziano G, Kittler H. Artificial intelligence in skin cancer diagnosis: a reality check. J Invest Dermatol. 2024;144(3):492–499. doi:10.1016/j.jid.2023.10.004",
 "tschandl2020": "Tschandl P, Rinner C, Apalla Z, Argenziano G, Codella N, Halpern A, et al. Human–computer collaboration for skin cancer recognition. Nat Med. 2020;26(8):1229–1234. doi:10.1038/s41591-020-0942-0",
 "chabi": "Chabi Adjobo E, Sanda Mahama AT, Gouton P, Tossa J. Automatic localization of five relevant dermoscopic structures based on YOLOv8 for diagnosis improvement. J Imaging. 2023;9(7):148. doi:10.3390/jimaging9070148",
 "shentu": "Shentu J, Watson M, Al Moubayed N. Controllable synthesis of dermoscopic images using diffusion models for enhanced computer aided diagnosis and detection. Med Image Anal. 2026;113:104191. doi:10.1016/j.media.2026.104191",
 "tobin": "Tobin J, Fong R, Ray A, Schneider J, Zaremba W, Abbeel P. Domain randomization for transferring deep neural networks from simulation to the real world. In: IEEE/RSJ International Conference on Intelligent Robots and Systems (IROS); 2017. p. 23–30. doi:10.1109/IROS.2017.8202133",
 "talavera": "Talavera-Martinez L, Tschandl P. Automated dermatoscopic pattern discovery by clustering neural network output for human-computer interaction. arXiv:2309.08533. 2023.",
}
ORDER = []


def cite(m):
    keys = [k.strip().lstrip("@") for k in m.group(1).split(",")]
    nums = []
    for k in keys:
        if k not in REFS: raise KeyError(k)
        if k not in ORDER: ORDER.append(k)
        nums.append(ORDER.index(k) + 1)
    nums = sorted(set(nums)); parts = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        parts.append(str(nums[i]) if j == i else (f"{nums[i]},{nums[j]}" if j == i + 1 else f"{nums[i]}–{nums[j]}")); i = j + 1
    return "[" + ",".join(parts) + "]"


# ------------------------------------------------------------------ text helpers
TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def runs(par, text, size=None, bold=False, italic=False):
    text = re.sub(r"\[(@[^\]]+)\]", cite, text)
    for part in TOK.split(text):
        if not part: continue
        b, i = bold, italic
        if part.startswith("**"): part, b = part[2:-2], True
        elif part.startswith("*"): part, i = part[1:-1], True
        r = par.add_run(part); r.bold, r.italic = b, i
        if size: r.font.size = Pt(size)
    return par


def P(text, size=None, align=None, after=None, before=None, indent=None, keep=False, style=None):
    p = doc.add_paragraph(style=style); runs(p, text, size)
    f = p.paragraph_format
    if align is not None: f.alignment = align
    if after is not None: f.space_after = Pt(after)
    if before is not None: f.space_before = Pt(before)
    if indent is not None: f.left_indent = Cm(indent)
    if keep: f.keep_with_next = True
    return p


def H(text, level=1):
    return doc.add_heading(text, level)


def bullets(items, numbered=False, size=None):
    for k, it in enumerate(items):
        p = doc.add_paragraph(); f = p.paragraph_format
        f.left_indent, f.first_line_indent, f.space_after = Cm(0.9), Cm(-0.5), Pt(2.5)
        mark = f"{k + 1}." if numbered else "•"
        runs(p, f"{mark}\t{it}", size)
        f.tab_stops.add_tab_stop(Cm(0.9))
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, sz in edges.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single" if sz else "nil"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000"); b.append(e)


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr; m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def caption(label, text, size=9):
    p = doc.add_paragraph(); f = p.paragraph_format
    f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; f.space_after = Pt(10); f.space_before = Pt(3)
    r = p.add_run(label + " "); r.bold = True; r.font.size = Pt(size)
    runs(p, text, size)
    return p


def table(header, rows, widths, size=8.5, label=None, cap=None, note=None, align=None, img_cols=()):
    if cap:  # table captions go above the table
        c = caption(label, cap); c.paragraph_format.keep_with_next = True; c.paragraph_format.space_after = Pt(3)
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; cell_margins(t)
    for c_i, w in enumerate(widths):
        t.columns[c_i].width = Cm(w)
    tblW = t._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); t._tbl.tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 567))); tblW.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
    for r_i, row in enumerate([header] + rows):
        tr = t.rows[r_i]
        if r_i == 0:
            trPr = tr._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
        trPr = tr._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        for c_i, val in enumerate(row):
            cell = tr.cells[c_i]; cell.width = Cm(widths[c_i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; f = p.paragraph_format; f.space_after = Pt(0); f.line_spacing = 1.0
            f.alignment = (align[c_i] if align else WD_ALIGN_PARAGRAPH.LEFT)
            if c_i in img_cols and r_i > 0 and val:
                path, w = val; p.add_run().add_picture(path, width=Cm(w)); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                runs(p, str(val), size, bold=(r_i == 0))
            e = dict(top=0, bottom=0, left=0, right=0)
            if r_i == 0: e.update(top=12, bottom=6)
            if r_i == len(rows): e.update(bottom=12)
            border(cell, **e)
    if note:
        n = P(note, size=8, after=10, before=3)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figure(path, width_cm, label, cap):
    p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(label, cap)


L, C, R = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT

# ------------------------------------------------------------------ header / footer
def field(par, instr):
    r = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r3._r.append(f2)


sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]; hp.alignment = R
r = hp.add_run("Synthetic glyphs for dermoscopic structure detection · Preprint v8"); r.italic = True; r.font.size = Pt(8.5)
for ft in (sec.footer, sec.first_page_footer):
    fp = ft.paragraphs[0]; fp.alignment = C; field(fp, "PAGE")

# ================================================================== FRONT MATTER
p = P("PREPRINT · NOT PEER REVIEWED", size=8.5, align=L, after=10); p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
P("**Synthetic Glyphs for Dermatologist-Oriented Structure Detection: An Explainable Second Machine-Learning Opinion on Dermoscopic Lesions**",
  size=17, align=L, after=10).paragraph_format.line_spacing = 1.05
P("F. Wolz", size=11.5, align=L, after=1)
P("*Independent researcher, Germany*", size=9.5, align=L, after=1)
P("Correspondence: via the project repository, https://github.com/FaGit99/Glyphs (GitHub issues)", size=9, align=L, after=1)
P("Version 8 · 22 October 2026", size=9, align=L, after=12)


def abstract_block():
    rule = doc.add_paragraph(); rule.paragraph_format.space_after = Pt(2)
    pPr = rule._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); e = OxmlElement("w:top")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000"); bd.append(e); pPr.append(bd)
    P("**Abstract**", size=11, align=L, after=4)
    items = [
     ("Background", "Deep-learning classifiers for dermoscopy reach high diagnostic accuracy but do not explain their output in the structural vocabulary dermatologists use. Structure-level supervision could close this gap, yet public pixel-level masks exist for only five dermoscopic structures."),
     ("Objective", "To build an explainable, dermatologist-oriented second opinion that names, localises and qualifies dermoscopic structures, and to quantify what a detector trained on synthetic structure renderings (“glyphs”) achieves on its own and when combined with dermatologist-annotated images."),
     ("Methods", "We implemented 71 of the 84 terms of the International Dermoscopy Society terminology consensus as parametric glyph generators and composited them onto real lesion-free skin, yielding scenes with exact region labels. A ViT-S/16 detector with an image head and a region head was trained under seven regimes: synthetic only, real masks only, and mixtures of synthetic data with ISIC 2018 masks and 248 expert-annotated EASY images, including dose experiments that varied the number of masks per structure with and without synthetic data. Evaluation used multi-reader expert masks (EASY, four-fold out-of-fold), ISIC 2018 Task 2, the Derm7pt 7-point checklist and four external melanoma-versus-nevus datasets."),
     ("Results", "Trained on synthetic glyphs alone, the detector recognised expert-annotated structures well above chance (mean AUROC 0.72, 95% CI 0.71–0.74; up to 0.92 for blue-whitish veil), and its maps stayed within the lesion. Adding ISIC 2018 masks raised the mean AUROC to 0.79, and adding dermatologist masks raised it to 0.90 (0.88–0.91). Per structure, 3–4 masked images gave an AUROC of 0.80 and 4–7 gave 0.85 for the two most responsive structures; once two or more masks were available, synthetic pre-training no longer improved recognition, even with equal loss weighting. Real-only training matched the mixture on annotated structures (0.90) but fell below chance on structures without masks (0.41 vs 0.70)."),
     ("Conclusions", "Synthetic glyphs yield working detectors across the consensus vocabulary without any real labels, and a handful of dermatologist masks per structure then brings recognition close to its ceiling. From that point, synthesis contributes coverage of unmasked structures rather than accuracy. The resulting structure analysis is shown alongside a direct classifier, together with the validation status of every structure."),
    ]
    for head, txt in items:
        P(f"**{head}.** {txt}", size=9.8, after=3.5)
    P("**Keywords:** dermoscopy; synthetic training data; explainable artificial intelligence; vision transformer; melanoma; dermoscopic structures; second opinion",
      size=9.8, after=2, before=4, align=L)
    rule = doc.add_paragraph(); pPr = rule._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000"); bd.append(e); pPr.append(bd)
    rule.paragraph_format.space_after = Pt(6)


abstract_block()

# ================================================================== 1 INTRODUCTION
H("1  Introduction")
P("Dermatologists reach a diagnosis by recognising patterns in skin lesions; a machine-learning second opinion intended for them should therefore express its findings in terms of those patterns. Direct classifiers and foundation models for dermatology now discriminate malignant from benign lesions with high accuracy [@panderm]; in this project, an EVA-02 base model [@eva] at 448 px reaches an AUROC of 0.9675 on its held-out test split. Such a score does not, however, indicate *why* a lesion is suspicious, and the clinical benefit of artificial intelligence in dermatology depends on how well its output can be combined with human expertise [@tschandl2020,@brancaccio]. Dermatologists describe pigment network, streaks, globules, veils and vessel patterns, and they base their decisions on the combination and location of these structures. This vocabulary is standardised in the terminology consensus of the International Dermoscopy Society (IDS) [@kittler] and underlies diagnostic algorithms such as the 7-point checklist [@argenziano].")
P("The obstacle to a structure-level model is data. Public pixel-level structure masks exist for five attributes in ISIC 2018 Task 2 [@codella]: pigment network, negative network, streaks, milia-like cysts and globules. Detectors trained on these masks, including recent object-detection approaches [@chabi], are therefore limited to these five structures. The EASY expert-agreement study [@liopyris] provides multi-reader masks for 31 melanocytic features, but on only 248 exemplar images, and Derm7pt [@kawahara] provides image-level labels for the 7-point criteria. Most of the 84 consensus terms have no real annotation at all. EASY also shows that expert readers agree on the location of only about one fifth of the marked area [@liopyris], so real masks are both scarce and noisy.")
P("Synthetic data are an established route to labels at scale. In dermoscopy, diffusion models have been used to synthesise whole lesion images to augment diagnostic classifiers [@shentu]; such images carry class labels but no structure-level masks. In robotics, domain randomisation has shown that models trained solely on procedurally rendered scenes can transfer to real images [@tobin]. Unsupervised pattern discovery offers a further route to interpretable output [@talavera], but it does not use the consensus vocabulary.")
P("This paper takes a procedural route. Each consensus term is drawn as a parametric *glyph*, glyphs are composited onto real lesion-free skin, and the resulting scenes carry exact region labels at no annotation cost. We address four questions:")
bullets(["What can a detector trained on synthetic glyphs alone recognise and localise on real images?",
         "What does mixing in real, dermatologist-masked lesions add, for which structures, and how many masks are sufficient?",
         "Does synthesis still contribute once real masks are available?",
         "How can detected structures be presented to dermatologists as a second opinion next to a direct classifier, without presenting unvalidated output as established?"], numbered=True)
P("Our contributions are (i) a glyph register of 71 drawable IDS structures with generators, parameters and training frequencies; (ii) a controlled comparison of seven training regimes on expert-annotated, external and checklist data; (iii) dose experiments that measure how many dermatologist masks a structure needs, with and without synthetic pre-training; and (iv) a two-path report that places the structure analysis, with an explicit validation status per structure, next to the direct classifier.")

# ================================================================== 2 GLYPHS
H("2  Glyph taxonomy and synthetic scene generation")
H("2.1  Taxonomy", 2)
P("The taxonomy has 84 entries (DS001–DS084) derived from Tables I–III of the IDS consensus [@kittler]. Each entry holds a metaphoric term (e.g. “blue-whitish veil”) and a descriptive term (“structureless zone, blue”); a category (lines, clods, dots, circles, structureless, vessels or patterns); site and imaging constraints (facial, acral/volar, polarised only); and the diagnoses the consensus associates with the structure. Thirteen entries are not drawn: three are modifiers (regular/irregular; DS036, DS044, DS058) and ten describe dataset criteria or colours rather than shapes (DS075–DS084). Seventy-one terms are therefore rendered, by 21 generator families.")
H("2.2  Generators and scene composition", 2)
P("A registry maps each ID to a generator, its parameters, a parent label (e.g. network), a drawing layer, and site and polarisation flags. The largest families are vessels (12 IDs, including dotted, glomerular, hairpin, linear irregular, corkscrew, comma, arborizing and polymorphous vessels), clods and globules (8), network (7), dots (7), structureless zones (6), streaks (5) and acral parallel patterns (4).")
P("A scene starts from a real lesion-free skin window from HAM10000 [@ham] at native scale (224 px). The lesion itself is synthetic: its core and rim colours are sampled relative to the skin tone from 1,341 real HAM10000 lesions. Glyphs are drawn at 448 px in a fixed layer order (zones, network, clods, dots, streaks, vessels) and blended at the native pixel pitch. Hair and air bubbles are then added as unlabelled distractors. Figure 1 shows one elementary unit per drawn ID.")
P("Two early pipeline errors that distorted realism measurements were identified and removed. First, JPEG re-encoding created a detectable artefact on its own: a real-versus-synthetic classifier on skin-only windows reached an AUC of 0.94 with re-encoding and 0.59 without. Second, lesion size was mismatched: real native crops are filled to a median of 88% by the lesion, whereas the first synthetic scenes were filled to only 30%. A hand-crafted real-versus-synthetic classifier still separates the final scenes with an AUC of about 0.95, dominated by edge sharpness. Realism was deliberately not optimised further against this metric; the criterion that matters is transfer to real images (Section 4).")
figure(f"{G}/out/register/overview.png", 14.5, "Figure 1.",
       "Glyph register overview. One isolated elementary unit per drawn taxonomy ID (DS001–DS074), rendered by the training generator on a neutral background. In training scenes the units are repeated, coloured relative to real skin and lesion tones, and blended at the native pixel pitch.")
H("2.3  Labels", 2)
P("Each scene stores the visible IDs and one region mask per ID on a 28 × 28 grid. A region is the union of all instances of an ID, closed morphologically with a 21-px kernel so that scattered globules form one region. The training set has 16,000 scenes and the validation set 1,600. A scene carries 2.8 labels on average, and about one seventh of the scenes are structure-free. Each of the 67 non-acral IDs appears in 3.8–4.5% of training scenes. The four acral parallel patterns (DS015–DS018) have a working generator but are not part of the current training set.")
H("2.4  Glyph register", 2)
P("Table 1 lists 20 of the 71 register entries together with their real-data performance. The full register, with an element image and a training scene for every ID, is given in Supplementary Table S1.")
T1 = [
 ["DS002", "Delicate network", "Lines, reticular and thin", "g_network", "area", "4.0", "0.41 → 0.79"],
 ["DS004", "Atypical pigment network", "Lines, reticular and thick, or varying in colour", "g_network", "area", "4.2", "0.77 → 0.92"],
 ["DS005", "Negative pigment network", "Lines, reticular, hypopigmented, around brown clods", "g_network", "area", "4.0", "0.78 → 0.94"],
 ["DS006", "Shiny white streaks", "Lines, white, perpendicular", "g_white_lines", "3 / 4.4 / 7", "4.1", "0.68 → 0.93"],
 ["DS009", "Radial streaming", "Lines, radial and segmental", "g_streaks", "1 / 5.9 / 14", "3.8", "0.79 → 0.97"],
 ["DS023", "Starburst pattern", "Pseudopods or radial lines, circumferential", "g_streaks", "12 / 58 / 108", "4.0", "0.91 → 1.00"],
 ["DS024", "Pseudopods", "Streaks with bulbous ends", "g_streaks", "0 / 20.8 / 69", "4.1", "0.86 → 0.95"],
 ["DS025", "Globules", "Clods, small, round or oval", "g_blobs", "13 / 23 / 35", "4.1", "0.50 → 0.89 (regular)"],
 ["DS026", "Rim of brown globules", "Clods, brown, circumferential", "g_blobs", "1 / 1.7 / 3", "4.0", "0.65 → 0.99"],
 ["DS029", "Cobblestone pattern", "Clods, large and polygonal", "g_cobblestone", "area", "4.2", "0.53 → 0.92"],
 ["DS033", "Milky-red globules", "Clods, pink and small", "g_blobs", "5 / 10 / 16", "4.1", "— (n+ = 7)"],
 ["DS038", "Peppering", "Dots, grey", "g_dots", "44 / 129 / 187", "4.1", "0.68 → 0.92"],
 ["DS048", "Blotch", "Structureless zone, brown or black", "g_zone", "1 / 1.1 / 2", "4.5", "0.85 → 0.86 (irregular)"],
 ["DS049", "Blue-whitish veil", "Structureless zone, blue", "g_zone", "1", "4.5", "0.92 → 0.97"],
 ["DS050", "Milky-red areas", "Structureless zone, pink", "g_zone", "1", "4.5", "0.74 → 0.94"],
 ["DS060", "Dotted vessels", "Vessels, dots", "g_vessel", "4 / 9.9 / 16", "4.1", "0.74 → 0.90"],
 ["DS065", "Linear irregular vessels", "Vessels, serpentine", "g_vessel", "1 / 1.1 / 2", "4.4", "0.85 → 0.96"],
 ["DS067", "Comma vessels", "Vessels, curved", "g_vessel", "2 / 2.8 / 5", "4.2", "0.87 → 0.96"],
 ["DS069", "Polymorphous vessels", "Vessels, polymorphous (several types)", "g_vessel_set", "7 / 11.7 / 25", "4.4", "0.89 → 0.97"],
 ["DS072", "Arborizing vessels", "Vessels, branched", "g_vessel", "1 / 1.1 / 2", "3.9", "0.66 → 0.88 (Derm7pt, n+ = 13)"],
]
table(["ID", "Metaphoric term", "Descriptive term (IDS)", "Generator", "Units per scene", "Scenes (%)", "AUROC R0 → R3"], T1,
      [1.25, 2.6, 4.1, 2.1, 2.0, 1.35, 2.8], size=8, label="Table 1.", align=[L, L, L, L, C, C, L],
      cap="Extract of the glyph register (20 of 71 drawn IDs).",
      note="Units per scene: minimum / mean / maximum number of elementary units; “area” denotes an area texture. Scenes: share of the 16,000 training scenes that contain the ID. AUROC R0 → R3: image-level AUROC of the mapped ID channel on EASY for synthetic-only training (R0) and for the mixed model (R3, out-of-fold); qualifiers in parentheses name the EASY feature used. For arborizing vessels, which have no EASY target, the Derm7pt zero-shot AUROC is given. —, no public real target with ≥10 positive images; n+, number of positive images.")

# ================================================================== 3 METHODS
H("3  Methods")
P("One detector architecture was trained under seven data regimes and evaluated with the same protocol on expert-annotated, checklist and external data.")
H("3.1  Detector", 2)
P("The detector is a Vision Transformer ViT-S/16 [@vit] pre-trained on ImageNet-21k and implemented with the timm library [@timm]. Two heads share the backbone: an image head on the CLS token with one sigmoid output per channel, and a region head that outputs a 14 × 14 probability map per channel. The output comprises 76 channels: 71 glyph-ID channels and 5 channels for the ISIC 2018 Task 2 attributes (pigment network, negative network, streaks, milia-like cysts and globules). In synthetic scenes, each attribute channel is filled with the union of its mapped IDs. The loss is binary cross-entropy, with the map loss weighted ×5 and a positive-class weight of 3; rotations and flips are used as augmentation. Real images supervise only the channels for which they carry a target; all other channels are masked out of the loss.")
H("3.2  Real data and their roles", 2)
P("Table 2 summarises the real datasets. Splits use blocks of consecutive ISIC IDs, which keeps images of the same lesion on one side of the split.", keep=True)
table(["Dataset", "Size", "Role", "Split"], [
 ["HAM10000 [@ham]", "10,015 images", "Skin windows and lesion palette for synthesis; external test", "Index parity 0 for synthesis, parity 1 for tests"],
 ["ISIC 2018 Task 2 [@codella]", "2,594 images, 5 attribute masks", "Real masks (R1–R5)", "Blocks of 25 consecutive IDs: 1,556 / 519 / 500 (train / calibration / test)"],
 ["EASY [@liopyris]", "248 exemplar images, 31 features, 4–6 readers each", "Dermatologist masks (R3–R5) and expert test set", "4 folds by EASY group, out-of-fold"],
 ["Derm7pt [@kawahara]", "1,011 cases, 7-point labels", "Checklist criteria (image level)", "Official 413 / 203 / 395"],
 ["BCN20000 [@bcn], ISIC 2019 [@isicdata], ISIC 2020 [@isic2020]", "500 / 388 / 500 melanomas; 800 nevi each (also for HAM10000)", "External melanoma-vs-nevus test only", "ISIC 2019 restricted to images not in HAM10000, BCN20000 or ISIC 2018"],
], [3.6, 3.6, 4.2, 4.8], size=8.5, label="Table 2.", cap="Real datasets used in this study and their roles.")
P("**EASY targets.** EASY features were mapped manually to taxonomy IDs: 22 unambiguously and 9 coarsely (e.g. regular and irregular dots → DS037). A grid cell is positive when at least two readers marked the feature and negative when no reader did; cells marked by exactly one reader are ignored. Twenty-four ID channels receive EASY targets. Four coarsely mapped features (homogeneous, structureless brown, scar-like depigmentation and angulated lines) are not trained and are reported separately.")
H("3.3  Training regimes", 2)
P("Table 3 defines the seven regimes. All regimes share architecture, folds, random seed and number of optimiser steps (16 × 166). A source that a regime omits is simply dropped from the batch; in R5 the 47 ID channels without an EASY mapping therefore receive no supervision. Seed variance was checked with a second seed for R0 and R3.", keep=True)
table(["Regime", "Synthetic glyphs", "ISIC 2018 masks", "EASY dermatologist masks"], [
 ["R0  Synthetic only", "yes", "–", "–"],
 ["R1  Real only (ISIC)", "–", "yes (5 channels)", "–"],
 ["R2  Synthetic + ISIC (V2)", "yes", "yes", "–"],
 ["R3  Synthetic + ISIC + EASY (V3)", "yes", "yes", "yes, all training-fold images"],
 ["R4  Dose curve", "yes", "yes", "25 / 50 / 75% of training-fold images (46 / 93 / 140 of 186)"],
 ["R5  Real only (ISIC + EASY)", "–", "yes", "yes"],
 ["R6  Per-structure dose (arms A / B / A′)", "A, A′: yes; B: –", "yes", "all images; masks of 3 structures limited to n positives"],
], [5.0, 2.6, 2.8, 5.8], size=8.5, label="Table 3.", cap="Training regimes.", align=[L, C, C, L])
P("**Per-structure dose experiment (R6).** Because EASY features co-occur, subsampling whole images (R4) does not control how many masks a given structure receives. R6 therefore varies only the supervision of three well-represented structures: peppering (DS038), shiny white structures (DS006) and dots (DS037, which covers both regular and irregular dots). Every EASY image stays in training for all other channels. For each structure, n = 0, 2, 4, 8, 16 or 32 positive training images keep their mask, or all of them (39–64 per fold); the remaining positives are excluded from the loss for that channel only, and negatives are unchanged. Draws are independent per structure, nested across n within a seed, and identical across arms. Arm A uses synthetic glyphs, ISIC 2018 and EASY as in R3; arm B omits synthetic glyphs as in R5; both use three seeds. The loss is averaged over all supervised entries of a batch, so real masks carry a larger share of it when synthetic rows are absent. Arm A′ therefore repeats arm A with the loss averaged separately over synthetic and real rows, giving real masks the same weight as in arm B (two seeds; n = 0, 2, 4, 8 and all). The dose is reported as the number of positive images per fold that kept image-level supervision.")
H("3.4  Lesion zones and checklist channels", 2)
P("A lesion segmenter (ViT-S at 384 px with a small decoder) reaches a Dice coefficient of 0.907 on the 500 ISIC 2018 test images. Its mask defines three zones: lesion core, border (±16 px) and outside. Structures are diagnostic inside the core and border; detections outside serve as a plausibility check.")
P("For the deployed model, 14 checklist channels were fine-tuned on Derm7pt image labels by multiple-instance learning. Each image is a bag of up to 16 tiles covering at least 15% lesion, and the bag score is the mean of the top three tiles. The channels were then Platt-calibrated [@platt] on out-of-fold predictions. The regime comparison in Section 4 uses the frozen regime models without this fine-tuning, so that all regimes are treated equally.")
H("3.5  Evaluation protocol", 2)
P("Every image is resampled to the native pixel pitch of the training data (short side 450 px). On EASY, the protocol uses four corner crops, a centre crop and a whole-image view (224 px each); the image score is the maximum over views and cells. Six capabilities were measured:")
bullets([
 "**C1 Presence:** image-level AUROC.",
 "**C2 Specificity:** rank of the mapped channel among the 71 ID channels, and the AUROC gain over a generic score (the mean of all 71 ID scores).",
 "**C3 Localisation:** pointing game against a lesion-referenced chance level (the segmenter mask dilated by one cell); cell AUROC inside lesion and border; patch IoU on ISIC 2018.",
 "**C4 Plausibility:** share of the top region (≥50% of the map maximum) that lies outside lesion and border.",
 "**C5 Coverage:** number of IDs with AUROC ≥0.80 on any real target with ≥10 positive images.",
 "**C6 Downstream:** melanoma versus nevus, scored with the 7-point rule on the channel scores and with a logistic probe fitted on Derm7pt training and validation data.",
])
P("**Statistics.** 95% confidence intervals (CIs) were obtained from 1,000 bootstrap resamples over images, paired across regimes.")

# ================================================================== 4 RESULTS
H("4  Results")
P("Synthetic glyphs alone yielded a working, lesion-confined structure detector. Mixing in dermatologist masks raised the mean AUROC on expert-annotated images from 0.72 to 0.90; per structure, a handful of masks gave most of this gain. Table 4 summarises all capabilities by regime.")
H("4.1  Capabilities by training regime", 2)
T4 = [
 ["C1 EASY, mean AUROC (25 trained features)", "0.724 (0.71–0.74)", "–", "0.785 (0.77–0.80)", "**0.896 (0.88–0.91)**", "0.904 (0.89–0.91)"],
 ["C1 ISIC 2018 test, mean AUROC (5 attributes)", "0.620", "0.896", "0.883", "0.889", "0.893"],
 ["C1 Derm7pt zero-shot, mean of 9 criteria", "0.682", "–", "0.747", "**0.819**", "0.710"],
 ["C1 Derm7pt, criteria without EASY-supervised ID†", "0.596", "–", "0.654", "**0.697**", "0.410"],
 ["C2 Mapped channel ranked 1st / top 3 (of 71)", "20% / 40%", "–", "32% / 60%", "**80% / 96%**", "76% / 88%"],
 ["C2 AUROC gain over generic score", "+0.16", "–", "+0.21", "+0.31", "+0.32"],
 ["C3 EASY pointing game (lesion-referenced chance)", "0.64 (0.37)", "–", "0.72 (0.37)", "**0.80** (0.38)", "0.77 (0.38)"],
 ["C3 ISIC 2018 test, patch IoU (5 attributes)", "0.10", "0.32", "0.29", "0.32", "0.33"],
 ["C4 EASY top-region share outside lesion + border‡", "4.6%", "–", "0.8%", "**0.6%**", "3.0%"],
 ["C5 IDs with real-data AUROC ≥0.80 / evaluable", "8 / 29", "0 / 0", "19 / 29", "**24 / 29**", "21 / 29"],
 ["C6 Melanoma vs nevus, probe, Derm7pt test", "0.769", "0.635", "0.717", "0.844", "0.866"],
 ["C6 Same probe, 4 external sets pooled§", "0.610", "0.503", "0.540", "**0.686**", "0.669"],
]
table(["Capability and metric", "R0 synthetic", "R1 ISIC only", "R2 synth. + ISIC", "R3 synth. + ISIC + EASY", "R5 ISIC + EASY"], T4,
      [5.1, 2.25, 1.6, 2.25, 2.6, 2.4], size=8, label="Table 4.", align=[L, C, C, C, C, C],
      cap="Capabilities by training regime. Values are means with 95% bootstrap CIs in parentheses where computed; best mixed-regime value in bold.",
      note="†Streaks, regression and arborizing vessels. ‡The area outside lesion and border covers about 45% of the image. §1,888 melanomas and 3,200 nevi. C6 uses the frozen regime models without checklist fine-tuning; its values serve only to compare regimes. –, not applicable (no corresponding channel).")
H("4.2  What synthetic glyphs alone achieve (R0)", 2)
P("**Presence.** Without a single real label, the R0 detector recognised dermatologist-annotated structures well above chance: the mean AUROC on EASY was 0.724 (95% CI 0.71–0.74). Retraining with the final pipeline and a second seed gave 0.702 and 0.706. The strongest structures were those with a distinctive shape or colour: blue-whitish veil (0.92), starburst pattern (0.91), polymorphous vessels (0.89), comma vessels (0.87), pseudopods (0.86), and broadened network, irregular blotch and linear irregular vessels (0.85 each).")
P("**Specificity.** The mapped channel exceeded the generic all-channel score by +0.16 AUROC, so the detector does not merely measure how melanoma-like an image is. For vessels, the generic score reached only 0.19–0.45, whereas the mapped vessel channels reached 0.74–0.89.")
P("**Localisation.** Maps were plausible but coarse. The map maximum fell on an expert-marked cell in 64% of positive images, against 37% lesion-referenced chance. Only 4.6% of the top region lay outside lesion and border, although that zone covers about 45% of the image. Boundaries were poor: the ISIC 2018 patch IoU was 0.10, compared with 0.32 for models trained on real masks.")
P("**Failures.** R0 failed on structures whose real appearance depends on fine texture or subtle colour: delicate (0.41) and typical (0.62) pigment network, regular globules (0.50), cobblestone pattern (0.53), branched streaks (0.50) and shiny white streaks (0.68).")
P("**Coverage.** Synthesis provided a trained channel for 67 of the 71 drawn IDs (all except the four acral patterns). Only 29 of these can be checked against any public real data, and 8 of them reached an AUROC ≥0.80. The lower 95% confidence bound exceeded 0.5 for 19 of the 25 EASY-evaluable IDs.")
H("4.3  What mixing with real dermatologist masks adds (R2, R3)", 2)
P("Adding ISIC 2018 masks (R0 → R2) increased the mean EASY AUROC by +0.06 (95% CI +0.05 to +0.08). The gain was concentrated in the network and globule families covered by ISIC 2018: typical network +0.23, delicate network +0.23, regular globules +0.26 and rim of brown globules +0.26. ISIC masks also raised the ISIC 2018 IoU from 0.10 to 0.29.")
P("Adding EASY dermatologist masks (R2 → R3) contributed a further +0.11 (+0.10 to +0.13). The largest gains fell on structures absent from ISIC 2018: cobblestone pattern +0.34, shiny white streaks +0.34, peppering +0.30, milky-red areas +0.16, and linear irregular and comma vessels +0.15 each (Figure 2). With EASY, the mapped channel ranked first for 80% of features (R0: 20%), and the pointing-game hit rate rose to 0.80. Structures that were already strong changed little: blue-whitish veil, irregular blotch, broadened network and pseudopods.")
figure(f"{FIG}/fig2_features.png", 14.0, "Figure 2.",
       "Per-feature image-level AUROC on EASY for the 25 features trained in R3, sorted by the synthetic-only result (R0). Grey bars connect R0 and R3; the number of positive images (≥2 readers) is given in parentheses. The gap closes most where synthesis alone was weakest (network subtypes, globules, peppering, shiny white streaks); shape-defined structures (veil, starburst, vessels, pseudopods) were already strong with synthetic glyphs alone. The dashed line marks chance level.")
P("**Transfer beyond EASY.** The gains carried over to other data: the Derm7pt zero-shot criteria rose from 0.747 to 0.819, the frozen melanoma-versus-nevus probe on the Derm7pt test set from 0.717 to 0.844, and the same probe on four external sources from 0.540 to 0.686. ISIC 2018 localisation was unchanged or slightly better (IoU for negative network 0.22 → 0.31, streaks 0.15 → 0.20; Figure 3). The mixed model therefore improved where dermatologists annotated, without losing performance elsewhere.")
figure(f"{FIG}/fig3_regions.png", 16.0, "Figure 3.",
       "Region maps on held-out ISIC 2018 test images, two examples per attribute (rows: pigment network, negative network, streaks, milia-like cysts, globules). Within each example, left: ISIC ground-truth mask; middle: synthetic-only model (R0); right: mixed model (“hybrid”, R3). Yellow outlines mark regions above the calibrated threshold; white outlines mark the lesion. The synthetic-only maps respond but rarely exceed the threshold; the mixed model finds the annotated regions, although its boundaries remain coarse.")
H("4.4  How many dermatologist masks are sufficient? (R4, R6)", 2)
P("Subsampling whole EASY images (R4) gave a mean curve that rose steeply and then flattened (Figure 4). With 46 images per fold, the mean AUROC over the 25 trained features increased from 0.781 to 0.869 (+0.087, 95% CI +0.073 to +0.101); 93 images added a further +0.016, and beyond that gains fell within seed variance (R3 with a second seed: 0.889). This mean does not translate into a number of masks per structure. Because features co-occur, 93 images supplied each trained channel with a median of about 12 positive training images rather than three, and the mean combines structures that were still improving with others that did not respond.")
P("The per-structure experiment (R6, Figure 5) answers the question directly. With synthetic glyphs (arm A), 3–4 masked positive images raised peppering and shiny white structures to an AUROC of 0.80, 5–7 to 0.85 and 16–18 to 0.90; with all masks they reached 0.93 and 0.94. Irregular dots were harder: 0.80 required about 9 masks, and 64 masks reached only 0.85. Ninety per cent of the gain from zero to all masks was reached with 6 (95% CI 4–11), 13 (8–21) and 19 (10–33) masks, respectively. Peppering and shiny white structures were still rising at the largest dose. Regular dots (28 positives, sharing a channel with irregular dots) improved only from 0.58 to 0.69 and are not interpreted further. Recognition became specific early: with two masks the mapped channel ranked first or second of 71 for peppering and irregular dots, and with four masks first for shiny white structures. Localisation kept improving up to the full dose; for peppering the lesion-referenced pointing game rose from 0.36 to 0.82 (chance 0.26).")
P("Without synthetic glyphs (arm B), the channels started near chance (0.46–0.60 with no masks, against 0.59–0.70 in arm A) but caught up with two masks. From then on they were level with or slightly above arm A (differences −0.01 to +0.05; one of 18 comparisons with a 95% CI excluding zero). Arm B needed 3–4 masks for an AUROC of 0.80, 4 for 0.85 (16 for irregular dots) and 11–14 for 0.90 (not reached for irregular dots). Part of this edge reflects loss weighting: with the loss averaged per source (arm A′), the synthetic arm moved towards arm B, and the remaining differences between A′ and B at 2–8 masks (−0.04 to −0.02) were not significant. Synthetic pre-training thus provides a zero-shot starting point but does not reduce the number of masks needed once two or more are available. The one exception was regular dots, the structure with the fewest real positives, where A′ exceeded B at every dose of two masks or more (+0.05 to +0.11, all 95% CIs excluding zero).")
figure(f"{FIG}/fig4_dose.png", 14.0, "Figure 4.",
       "Image-subsample dose curve (R4). Mean image-level AUROC over the 25 trained EASY features as a function of the number of dermatologist-masked EASY images per training fold; synthetic glyphs and ISIC 2018 masks are always included. The point at 0 is a retrained synthetic + ISIC control with the same number of optimiser steps. The shaded band shows the 95% bootstrap CI; folds are defined by EASY group, with one subsample draw per point. Horizontal lines: synthetic-only (R0) and real-only (R5) references.")
figure(f"{FIG}/fig5_target_dose.png", 13.0, "Figure 5.",
       "Per-structure dose experiment (R6). Image-level AUROC on EASY (out-of-fold) for each target structure as a function of the number of its positive training images per fold that keep their dermatologist mask; all other EASY labels are unchanged. Arm A: synthetic glyphs + ISIC 2018 + EASY; arm B: without synthetic glyphs; arm A′: as A, with the loss averaged separately over synthetic and real rows so that real masks carry the same weight as in B. Lines show the mean over seeds (three for A and B, two for A′); bands show 95% CIs from a bootstrap over images and seeds. The rightmost point uses all masks (39–64 per fold). Dotted line: synthetic-only reference (R0). Regular and irregular dots share one channel (DS037) and therefore the same dose.")
H("4.5  Does synthesis still help once real masks exist? (R5)", 2)
P("Where dermatologist masks exist, it does not. Real-only training (R5) matched the mixture on the 25 EASY-trained features (0.904 vs 0.896; difference −0.008, 95% CI −0.016 to +0.002) and on ISIC 2018 (AUROC 0.893 vs 0.889; IoU 0.33 vs 0.32). Because real masks weigh more in the loss when synthetic rows are absent, this comparison is not fully balanced; the per-structure experiment with equal weighting (arm A′, Section 4.4) led to the same conclusion for structures with ample masks.")
P("Where no masks exist, synthesis is indispensable. On the Derm7pt criteria whose IDs have no EASY target (streaks, regression, arborizing vessels), R5 fell to 0.41, below chance, while the mixture held 0.70. On the four coarsely mapped EASY features, R5 reached 0.47 against 0.58. For the 42 IDs without any public real data, synthesis is the only source of supervision. Real-only training also let more of the top region spill outside the lesion (3.0% vs 0.6%). On external data the mixture was slightly more robust, outperforming R5 on the four external sources with the pooled 7-point rule (+0.019, 95% CI +0.009 to +0.031) and with the probe (+0.017, +0.001 to +0.031).")
P("In summary, the mixture retains the accuracy of real masks where they exist and the breadth of synthesis where they do not; neither source alone provides both.")
H("4.6  Deployed model and relation to the direct classifier", 2)
P("The deployed model is R3 with the Derm7pt checklist channels fine-tuned on top. Its 7-point rule on predicted probabilities reached an AUROC of 0.886 on the Derm7pt test set (individual checklist channels 0.86–0.97). On four external datasets, without any diagnosis training, it discriminated melanoma from nevus with the AUROCs shown in Table 5. Because the classic “≥3 points” threshold did not transfer to external data, the report displays the expected checklist score rather than a verdict.", keep=True)
table(["Dataset", "HAM10000", "ISIC 2019", "BCN20000", "ISIC 2020", "Pooled"],
      [["AUROC", "0.794", "0.774", "0.733", "0.707", "0.749"]],
      [2.6, 2.3, 2.3, 2.3, 2.3, 2.3], size=8.5, label="Table 5.", align=[L, C, C, C, C, C],
      cap="Melanoma-versus-nevus discrimination of the deployed model’s 7-point score on external datasets (no diagnosis training).")
P("The structure path is not an accuracy booster. On the direct classifier’s own test split (8,054 dermoscopic images, 1,766 malignant), the direct EVA-02 model alone reached an AUROC of 0.9675; a two-stage fusion with the glyph probe gave 0.9692 (+0.0016, 95% CI 0.0000 to +0.0035).")
P("Its value lies in explanation and in prompting a second look. Of the 5,554 test images that the direct model called benign, a disagreement flag (glyph score at or above the 90th percentile of benign validation images) marked 381. Of these, 7.1% were malignant, against a base rate of 1.9% (×3.7), and the flag caught 27 of the 105 malignant lesions missed by the direct model.")

# ================================================================== 5 REPORT
H("5  The dermatologist-oriented second opinion")
P("The report places the structure analysis next to the direct classifier, never merged into it, and labels every statement with the extent to which it has been validated on real data. It is generated as a self-contained HTML page per image in 0.2–0.5 s after model loading; two complete examples are shown in Supplementary Figures S1 and S2. From top to bottom, it contains five elements.")
P("**Two independent paths.** Path 1 is the direct classifier: its score, the decision threshold (fixed on validation data at 95% sensitivity) and the test sensitivity and specificity at that threshold. Path 2 is the structure analysis: the glyph score as a percentile of the benign reference and the expected 7-point checklist score. A state line names one of four states (both high, direct only, glyph flag, both low), each with its observed malignancy rate on the test split.")
P("**Image with layers.** Region contours per detected structure, the lesion outline and the border zone, with an opacity control.")
P("**Detected structures.** Each structure shows its zone distribution (core / border / outside) and the diagnoses the consensus associates with it [@kittler]. Calibrated probabilities are shown only where calibration on real data exists. Structures trained only on synthetic glyphs are collapsed by default and show a coarse level (high / medium / low), never a percentage.")
P("**7-point checklist (experimental).** For each criterion, the report gives its probability, its contribution in points and the contributing tiles as a toggle layer, explicitly marked as not a validated localisation. The “≥3 points” threshold is not presented as a verdict.")
P("**Diagnosis groups.** A fixed-order card lists, for each diagnostic group, the structures that were detected and the typical ones that were not. It is informational only and does not rank diagnoses: an equal-weight mapping from structures to diagnoses was near chance on Derm7pt (AUROC 0.46–0.61), because the consensus states where a structure occurs, not how specific it is.")
P("**Validation status and warnings.** Every structure carries one of four badges: region-validated (ISIC 2018 masks), expert-validated at image level (EASY AUROC), image level only (Derm7pt AUROC), or not validated. The report warns when more than 25% of a detection lies outside lesion and border, when a structure is plausible only at another body site (e.g. facial or acral), and when it can be judged only under polarised dermoscopy. The mixed model adds an expert-validated badge to 24 ID channels; previously only the 5 ISIC attributes and the 14 checklist criteria had real-data validation. This is the practical effect of the dermatologist masks: more structures can be shown with real-data validation rather than as synthetic-only hints.")

# ================================================================== 6 DISCUSSION
H("6  Discussion")
P("Synthetic glyphs and dermatologist masks do different jobs: synthesis provides breadth across the consensus vocabulary, and dermatologist masks calibrate it to how structures actually appear in real images.")
P("**Why synthesis works where it works.** The structures R0 recognised best are defined by geometry or a distinctive colour: veil, starburst, pseudopods, vessel shapes and blotches. A parametric drawing can capture these, and the descriptive IDS terminology [@kittler] is in effect a geometric specification. Synthesis falls short where real appearance depends on fine texture and subtle contrast: delicate and typical network, regular globules, cobblestone pattern and shiny white streaks. The remaining real-versus-synthetic gap (AUC ≈0.95, dominated by edge sharpness) points in the same direction.")
P("**Why a handful of dermatologist masks goes far, and why synthesis does not shorten it.** For well-represented structures, 3–4 masked images gave an AUROC of 0.80 and 4–7 gave 0.85; subtle structures such as irregular dots needed considerably more. We expected synthetic pre-training to reduce this number by placing each structure in its own channel beforehand. It did not: without synthesis the channels started near chance but caught up after two masks, also when real masks were weighted equally. A plausible explanation is that the ImageNet-pretrained backbone, together with real supervision of other channels on the same images, already provides features from which a few masks suffice to define a structure. Synthetic glyphs therefore supply a zero-shot starting point and coverage of structures without masks, not a shortcut to learning from masks. Real masks also sharpen *which* channel fires: with EASY, the mapped channel became the top channel for 80% of features.")
P("**Why the mixture, and not real data alone.** On structures that have masks, real-only training is as good as the mixture, so synthesis is no substitute for annotation. The mixture is nevertheless the appropriate design for a dermatologist-facing tool, for three reasons. Coverage: real-only training leaves every unannotated structure without a working detector and falls below chance on Derm7pt for streaks, regression and arborizing vessels (0.41); of the 71 drawable IDs only 29 have any public real target and only 24 have EASY masks. Plausibility: real-only maps spill outside the lesion more often (3.0% vs 0.6% of the top region). Robustness: on external data the mixture is slightly ahead (+0.02 pooled).")
P("**A sufficient mixture.** The practical rule from this study is to retain synthetic glyphs for the full vocabulary and to add dermatologist masks per structure: about 5–10 masked images for an AUROC around 0.85 for distinctive structures, and more for subtle ones such as irregular dots. Synthesis keeps its value for structures without masks and for those with very few real positives, as for regular dots. Annotation effort is best directed at structures not yet covered rather than at more images of covered ones. Because expert readers agree spatially on only about 20% of marked superpixels [@liopyris], the consensus rule used here (≥2 readers positive, single-reader marks ignored) matters as much as the number of images.")
P("**A second opinion, not a second classifier.** The structure path does not make the direct classifier more accurate (fusion +0.0016 AUROC). Its role is different: it tells the dermatologist which structures are present, where they are and how well each detector is validated, and it raises a “second look” flag when the two paths disagree. This matches the collaborative use of decision support that has been shown to benefit clinicians [@tschandl2020] and the call for AI that supports rather than replaces dermatological judgement [@brancaccio].")
P("**Future work.** The next step is a structured review with dermatologists in two parts: whether the glyphs are drawn correctly (fidelity to the consensus definitions, recognisability without labels, plausibility of colour, scale and placement), and whether the dermatologist report is understandable and useful in practice (comprehension of badges and zones, agreement with the shown structures, effect on attention and decisions, usability). Known issues to be confirmed include the rendering of scar-like depigmentation (DS051), the composition of polymorphous vessels (DS069) and the density of the vessel and dot families. The masks dermatologists draw or correct during this review are exactly the kind of “sufficient mixture” this study shows to be needed. The full protocol is given in Supplementary Section S3.")

# ================================================================== 7 LIMITATIONS
H("7  Limitations")
P("The central estimates are robust to the choice of seed, but several properties of the data may make the out-of-fold gains optimistic.")
P("**Fold design.** EASY folds follow the dataset’s group variable and are not split by patient or lesion, so out-of-fold gains for R3–R6 may be optimistic. The bias affects R3 and R5 equally; the comparison between them is therefore less affected than that between R3 and R0.")
P("**Exemplar selection.** EASY images are exemplars, mostly melanomas and nevi, in which features co-occur. A high AUROC can partly reflect lesion type; channel rank and the generic-score control are therefore reported alongside AUROC.")
P("**Small samples.** Several features have only 11–18 positive images (rim of brown globules, comma vessels, starburst pattern, cobblestone pattern, regular blotch, homogeneous pattern, branched streaks), and arborizing vessels have 13 positives in Derm7pt. Bootstrap intervals over images are probably too narrow.")
P("**Mappings.** Mappings between EASY or Derm7pt features and taxonomy IDs were made by hand. Nine EASY features are mapped only coarsely, and regression (DS051) is a poor match in every regime.")
P("**Seeds and subsamples.** Each dose point uses one subsample draw, and seed variance was checked only for R0 and R3. It is small for the mean over trained features (≤0.007) but large for the four untrained features (0.58 vs 0.63), so no conclusion is drawn from the latter. The per-structure experiment (R6) covers only three structures with 59–101 positives, uses three seeds (two for arm A′) and is capped at 39–64 masks per fold. Peppering and shiny white structures were still improving at that cap, so the number of masks needed to approach their ceiling may be underestimated. Rare or subtle structures may need more masks than the structures studied here.")
P("**Unvalidated structures and localisation.** Forty-two of the 71 IDs have no public real target; their channels are trained but unvalidated, and the report labels them as such. Region maps are validated against real masks only for the 5 ISIC attributes and the 24 EASY channels; checklist channels are validated at image level only.")
P("**Scope.** The study covers dermoscopic images only. Thresholds were fixed on the direct classifier’s validation split, the direct score is not calibrated, and performance across skin tones was not assessed.")
P("**No clinical evaluation yet.** Neither the glyph renderings nor the report have so far been reviewed by dermatologists, and no reader study has measured an effect on diagnostic decisions.")

# ================================================================== 8 CONCLUSIONS
H("8  Conclusions")
P("Rendering 71 IDS consensus structures onto real skin, with no real labels, yields a detector that recognises dermatologist-annotated structures on independent expert images (mean AUROC 0.72, up to 0.92 for blue-whitish veil), keeps its maps inside the lesion and provides a trained channel for 67 structures of the vocabulary, far beyond the five for which public masks exist. Mixed with dermatologist-masked lesions, the detector works considerably better: the mean AUROC rises to 0.90, the correct channel becomes the top-ranked one for 80% of structures, and the gains carry over to the 7-point checklist and to external datasets. Per structure, a handful of masks suffices for distinctive structures (3–4 masked images for an AUROC of 0.80, 4–7 for 0.85), whereas subtle structures need more. Once two or more masks exist, synthetic pre-training no longer improves recognition, even with equal loss weighting; its value lies in the zero-shot starting point and in structures with no or very few masks. Where masks exist, real data supply precision; where they do not, synthetic glyphs remain the only working detectors. The resulting dermatologist-oriented second opinion names the structures in a lesion, shows where they are and states how well each one is validated, next to, but separate from, the direct classifier.")

# ================================================================== DECLARATIONS
H("Declarations")
for head, txt in [
 ("Funding", "This work received no specific funding."),
 ("Competing interests", "The author declares no competing interests."),
 ("Ethics statement", "This study used only publicly available, de-identified image datasets released for research. No new patient data were collected, and no ethics approval was required. All datasets were used in accordance with their respective licences."),
 ("Data availability", "All real datasets are publicly available: HAM10000 (doi:10.7910/DVN/DBW86T); ISIC 2018, ISIC 2019, ISIC 2020 and BCN20000 via the ISIC Archive (https://challenge.isic-archive.com/data); Derm7pt (https://derm.cs.sfu.ca); and the EASY expert-annotation data (https://github.com/ISIC-Research/expert-annotation-agreement-data; ISIC Archive collection 166). The synthetic scenes can be regenerated with the released code using fixed seeds."),
 ("Code availability", "The code for glyph generation, training, evaluation and the report is available at https://github.com/FaGit99/Glyphs. Supplementary Table S2 maps every result in this paper to the script and output file that produce it."),
 ("Author contributions", "Wolz F: idea, conceptualisation, methodology, supervision, validation, and writing – review and editing."),
 ("Use of AI tools", "An AI assistant (Claude, Anthropic) was used throughout this work under the author’s direction. It wrote and ran the code for data generation, training and evaluation, performed the statistical analyses, produced the figures and drafted the manuscript text. The author defined the research questions and experiments, reviewed the results and the text, and takes full responsibility for the content of this article."),
 ("Acknowledgements", "The author thanks the creators of HAM10000, the ISIC Archive and challenge datasets, BCN20000, Derm7pt and the EASY expert-agreement study for making their data publicly available."),
]:
    P(f"**{head}.** {txt}", size=10)

# ================================================================== REFERENCES (placeholder; filled after supplement so all citations are known)
H("References")
ref_anchor = doc.add_paragraph()

# ================================================================== SUPPLEMENT
page_break()
P("**Supplementary Material**", size=14, align=L, after=2)
P("Synthetic Glyphs for Dermatologist-Oriented Structure Detection — Preprint v8", size=9.5, align=L, after=10)
H("S1  Full glyph register")
P("Supplementary Table S1 lists all 71 drawn IDs. *Element*: the training generator on a flat neutral background with the count forced to one unit (area structures: a central texture patch). *Training scene*: a training-style composition that contains only this ID. *Units per scene*: minimum / mean / maximum over 12 test scenes, given as label instances written by the generator and, after the dot, as connected components of the instance mask (approximate). *Scenes*: number and share of the 16,000 training scenes that contain the ID. Site restrictions are given with the layer. The four acral patterns (DS015–DS018) are drawn but not yet part of the training set.")
rows = []
with open(f"{G}/out/register/register.csv", newline="") as fh:
    for r in csv.DictReader(fh):
        term = r["metaphoric_term"] or "(no metaphoric term)"
        gen = f"{r['generator']}, layer {r['layer']}" + (f", site: {r['site']}" if r["site"] else "")
        par = r["params"].replace(";", "; ")
        rows.append([(f"{G}/out/register/elements/{r['ID']}.png", 1.5), (f"{G}/out/register/scenes/{r['ID']}.png", 2.9),
                     f"**{r['ID']}** {term}\n*{r['descriptive_term']}*", gen + (f"\n{par}" if par else ""),
                     f"{r['label_instances_per_scene']} · {r['connected_units_per_scene']}",
                     f"{r['train_images_with_id']} ({r['train_share_pct']}%)"])
table(["Element", "Training scene (ID alone)", "ID and term", "Generator and parameters", "Units per scene", "Scenes"], rows,
      [1.8, 3.2, 3.6, 4.1, 2.0, 1.5], size=7, label="Table S1.", img_cols=(0, 1), align=[C, C, L, L, C, C],
      cap="Full glyph register (71 drawn IDs) with element and scene images.")
sp = doc.paragraphs[-1]; sp._p.getparent().remove(sp._p)          # drop the spacer after Table S1 (it can spill onto an empty page)
H("S2  Report examples").paragraph_format.page_break_before = True
P("Two complete reports (Section 5) for held-out test images of the direct classifier, rendered from the self-contained HTML files. The interface is in German, the language of the intended reviewers. Layer toggles and opacity are interactive in the HTML file; the screenshots show the default state and are split into two halves (left: top of the page; right: continuation).")


def two_up(a, b, label, cap):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col in t.columns: col.width = Cm(8.0)
    for c, path in zip(t.rows[0].cells, (a, b)):
        c.width = Cm(8.0); pp = c.paragraphs[0]; pp.alignment = C; pp.add_run().add_picture(path, width=Cm(7.7))
    caption(label, cap)


two_up(f"{V5P}/dv1_0.png", f"{V5P}/dv1_1.png", "Figure S1.",
       "Example 1: disagreement flag (malignant lesion). The direct classifier stays below its threshold, the structure analysis lies above the 91st percentile of the benign reference, and the report raises the “second look” flag (state line, top). Blue-whitish veil and blotch are shown with their EASY validation badge; angulated lines are marked as not validated and carry a body-site warning. Image: ISIC_0064677.")
page_break()
two_up(f"{V5P}/dv2_0.png", f"{V5P}/dv2_1.png", "Figure S2.",
       "Example 2: both paths high (malignant lesion). Both paths agree; the structure list, the experimental 7-point checklist and the diagnosis-group card show what supports the result. Image: ISIC_3384621.")

page_break()
H("S3  Protocol for the planned dermatologist review")
P("The review has two parts: whether the glyphs are drawn correctly (Part A), and whether the report is useful (Part B).")
H("Part A: Fidelity of the glyph generation", 2)
P("**Material.** For each of the 71 drawn IDs: the isolated element, two training scenes (the ID alone and within a full composition) and the IDS descriptive term, all available from the glyph register (Table S1).")
P("**Task.** Dermatologists rate each ID on three points: (a) whether the rendering matches the consensus definition (5-point scale); (b) whether the structure is recognisable without its label (blind naming of a random subset); and (c) whether colour, scale and density are plausible and placement (core, border, body site) matches clinical reality.")
P("**Known issues to confirm.** Scar-like depigmentation (DS051) does not resemble a scar; polymorphous vessels (DS069) should combine several vessel types within one lesion, following the dermoscopedia vessel terminology; and the number of elementary units per scene for the vessel and dot families (DS060–DS074) may be too low.")
P("**Outcome.** A list of IDs to redraw, prioritised by clinical relevance and by low real-data AUROC (Table 1), and a proposal of which structures to mask next (about 5–10 masked images per uncovered structure, more for subtle ones; Section 4.4).")
H("Part B: Usability of the report", 2)
P("**Material.** Reports for a balanced case set: clear melanomas and clear nevi, cases in which the two paths disagree, and cases with detections outside the lesion.")
P("**Task.** Dermatologists work through the cases with and without the structure path and assess (a) comprehension: whether validation badges, zones and checklist are read correctly; (b) agreement: whether the shown structures and regions match what they see (per structure: agree / partly / disagree); (c) usefulness: whether the display would change their attention or decision; and (d) trust and workload, using standard questionnaires such as the System Usability Scale.")
P("**Outcome.** Measured agreement between dermatologists and the detector per structure, which doubles as an additional expert validation set; concrete changes to the display; and a protocol for a later reader study that measures the diagnostic effect.")

H("S4  Reproducibility map").paragraph_format.page_break_before = True
P("All code is in the project folder glyphs/ (Python with PyTorch, timm, scikit-learn and OpenCV). Every number in this paper is produced by a script whose output is stored in models/*.json. Results of the earlier V2 model are frozen in glyphs/frozen_V2/. Data generation and evaluation use fixed seeds.", keep=True)
table(["Step", "Scripts", "Outputs"], [
 ["Glyphs and scenes", "geometry.py, structures.py, render.py, skin.py, compose.py, gen_dataset.py, build_register.py", "data/synth_*.npz, out/register/"],
 ["Detector, hybrid, mixed model (V3)", "dense.py, hybrid.py, hybrid3.py, easy_data.py", "models/hybrid_eval*.json, models/hybrid_calibration*.json"],
 ["EASY evaluation", "easy_eval.py, easy_v3_pred.py, easy_v3_compare.py", "models/easy_v3_compare.json"],
 ["Regime comparison R0–R5 (Tables 4, Figures 2, 4)", "mission_train.py, mission_pred.py, mission_ext.py, mission_compare.py, mission_figures.py", "models/mission_compare.json, MISSION_RESULTS.md"],
 ["Per-structure dose experiment R6 (Figure 5)", "mission_train.py (--tids, --tdose, --srcnorm), td_eval.py, Papers/build_v8/v8_figs.py", "models/td_dose.json, models/td_dose_A2.json"],
 ["Lesion segmentation", "lesion_seg.py, pipeline_eval.py", "models/pipeline_eval.json"],
 ["Checklist channels", "d7_data.py, mil.py, train_melanoma.py, calibrate_mel.py, eval_melanoma.py", "models/melanoma_eval.json, models/mel_calibration.json"],
 ["External tests, direct classifier, two paths (Table 5)", "external_check.py, direct_model.py, fusion_val_test.py, twopath_fit.py", "models/external_check.json, models/fusion_val_test.json, models/twopath.json"],
 ["Report", "analyze.py, check_integration.py", "out/twopath/*.html"],
], [4.2, 6.6, 5.4], size=8, label="Table S2.", cap="Scripts and output files behind each result.")

# ------------------------------------------------------------------ fill references at anchor
for k in ORDER:
    pass
anchor = ref_anchor
for n, k in enumerate(ORDER, 1):
    newp = OxmlElement("w:p"); anchor._p.addnext(newp)
    from docx.text.paragraph import Paragraph
    rp = Paragraph(newp, anchor._parent)
    f = rp.paragraph_format; f.left_indent, f.first_line_indent, f.space_after = Cm(0.8), Cm(-0.8), Pt(3)
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT; f.tab_stops.add_tab_stop(Cm(0.8))
    r = rp.add_run(f"{n}.\t{REFS[k]}"); r.font.size = Pt(9)
    anchor = rp
ref_anchor._p.getparent().remove(ref_anchor._p)
unused = [k for k in REFS if k not in ORDER]
print("refs:", len(ORDER), "unused:", unused)
doc.core_properties.title = "Synthetic Glyphs for Dermatologist-Oriented Structure Detection"
doc.core_properties.author = "F. Wolz"
doc.core_properties.subject = "Preprint v8"
doc.save(OUT); print("saved", OUT)
